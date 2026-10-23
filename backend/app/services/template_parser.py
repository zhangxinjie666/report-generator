"""模板解析：读取 docx 模板，得到封面字段位置、章节结构、表格行列。

为什么需要这一步：
    Word 文档在程序眼里只是一串"按顺序排列的段落和表格"，
    程序不能像人一样"看"到哪里是标题、哪里该填内容。
    所以先解析出一张"模板地图"——第几段是章节标题、封面哪几段有待填字段、
    每个表格几行几列。第 6 步组装报告时，report_builder 拿着这张地图精准填空。

本模板的两个现实特征（都是实测发现的，见 _dump_template.py 诊断脚本）：
    1. 封面标签里夹着空格，如"姓    名"、"班    级"，
       直接搜"姓名"永远搜不到 → 匹配前必须把所有空白字符删掉。
    2. 章节标题的"一、二、三"是 Word 自动编号，编号存在 docx 内部的
       numbering.xml 里，不在段落文字中 → 只能用已知标题清单精确匹配。

docx 的本质：
    一个 docx 文件其实是个 ZIP 压缩包，里面装着若干 XML 文件。
    python-docx 负责打开压缩包、解析 XML，把这些内容变成 Python 对象：
        Document  整份文档
        Paragraph 一个段落（正文里每一行文字）
        Run       段落里"格式一致的一小截文字"（比如"姓名："后面那条下划线）
        Table     表格，由 Row（行）和 Cell（单元格）组成

直接测试本文件：
    .\\venv\\Scripts\\python.exe -m app.services.template_parser ..\\assets\\报告模板.docx
"""

import json
import re
import sys
from pathlib import Path

from docx import Document


class TemplateParser:
    """解析报告模板 docx，产出"模板地图"。"""

    # 本模板的章节标题清单（正文标题文字本身不带编号，编号是 Word 自动加的）
    _SECTION_TITLES = (
        "实验目的",
        "主要仪器设备及软件",
        "实验原理",
        "实验电路图",
        "实验内容和实验结果",
        "结果分析",
        "实验小结",
        "附录",
    )

    # 封面字段标签。匹配时用"标签+冒号"做前缀，避免误伤正文标题
    _COVER_LABELS = (
        "课程名称", "实验名称", "学院", "班级",
        "学号", "姓名", "指导教师", "学期",
    )

    # 匹配所有空白字符：半角空格、全角空格（\u3000）、制表符、换行等
    _WS_RE = re.compile(r"[\s\u3000]+")

    # 正文大标题（"实验名称"单独成段、后面不跟冒号的，是正文标题而非封面字段）
    _BODY_TITLE = "实验名称"

    def parse(self, template_path: str | Path) -> dict:
        """解析模板，返回 JSON 可序列化的字典（"模板地图"）。"""
        # Document() 打开 docx（本质是解压并解析里面的 XML）
        doc = Document(str(template_path))
        return {
            "paragraph_total": len(doc.paragraphs),
            "table_total": len(doc.tables),
            "cover_fields": self._find_cover_fields(doc.paragraphs),
            "body_title_index": self._find_body_title(doc.paragraphs),
            "sections": self._find_sections(doc.paragraphs),
            "tables": self._describe_tables(doc.tables),
        }

    @classmethod
    def _norm(cls, text: str) -> str:
        """删除所有空白字符，用于匹配。

        例："姓    名：  张鑫杰" -> "姓名：张鑫杰"
        """
        return cls._WS_RE.sub("", text)

    def _find_cover_fields(self, paragraphs) -> list[dict]:
        """找出封面上"标签：值"形式的字段行。

        判断标准：段落去空白后以"标签："或"标签:"开头。
        这样正文大标题"实验名称"（后面没有冒号）不会被误判成封面字段。
        """
        fields = []
        for i, para in enumerate(paragraphs):
            raw = para.text
            norm = self._norm(raw)
            for label in self._COVER_LABELS:
                if norm.startswith(f"{label}：") or norm.startswith(f"{label}:"):
                    # 值 = 原文里第一个冒号之后的内容（保留原始空格无所谓，剥掉即可）
                    sep_idx = min(
                        (j for j in (raw.find("："), raw.find(":")) if j != -1),
                        default=-1,
                    )
                    current_value = raw[sep_idx + 1:].strip() if sep_idx != -1 else ""
                    fields.append({
                        "label": label,
                        "paragraph_index": i,
                        "current_value": current_value,
                        "has_fill_line": any(run.underline for run in para.runs),
                    })
                    break  # 一个段落只算一个字段
        return fields

    def _find_body_title(self, paragraphs) -> int | None:
        """找正文大标题"实验名称"（去空白后恰好等于这四个字）。"""
        for i, para in enumerate(paragraphs):
            if self._norm(para.text) == self._BODY_TITLE:
                return i
        return None

    def _find_sections(self, paragraphs) -> list[dict]:
        """找出章节标题段落（文字与已知标题清单精确相等）。"""
        sections = []
        for i, para in enumerate(paragraphs):
            text = self._norm(para.text)
            if text in self._SECTION_TITLES:
                sections.append({
                    "paragraph_index": i,
                    "title": text,
                    "style": para.style.name if para.style else "Normal",
                })
        return sections

    def _describe_tables(self, tables) -> list[dict]:
        """概括每个表格的结构：几行几列 + 表头预览。

        只取第一行（rows[0]）当表头预览——表头最能说明这张表是干嘛的，
        全部单元格都打出来没有意义，第 6 步填数时反正要按行列坐标逐格填。
        """
        result = []
        for ti, table in enumerate(tables):
            header: list[str] = []
            if len(table.rows) > 0:
                header = [cell.text.strip()[:20] for cell in table.rows[0].cells]
            result.append({
                "table_index": ti,
                "rows": len(table.rows),
                "cols": len(table.columns),
                "header_preview": header,
            })
        return result


if __name__ == "__main__":
    # 命令行快速验证：python -m app.services.template_parser <模板路径>
    if len(sys.argv) < 2:
        print("用法: python -m app.services.template_parser <模板.docx路径>")
        sys.exit(1)
    result = TemplateParser().parse(sys.argv[1])
    # ensure_ascii=False 让中文正常显示而不是 \uXXXX 转义
    print(json.dumps(result, ensure_ascii=False, indent=2))
